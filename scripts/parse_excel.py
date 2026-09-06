"""
Convierte el Excel de asistencia en data.json para el dashboard.
Detecta automáticamente:
  - Qué hojas existen en el libro (Ministerio, Regional, etc.)
  - Qué columnas de mes tienen encabezado válido (ENE, FEB, ..., DIC)
Así, cuando agregues datos de un nuevo mes (ej. SEPT) o una hoja nueva,
no hay que tocar este script: se detecta solo.
"""
import glob
import json
import os
import sys

import openpyxl

MESES_VALIDOS = ['ENE', 'FEB', 'MAR', 'ABR', 'MAY', 'JUN', 'JUL', 'AGO',
                  'SEPT', 'OCT', 'NOV', 'DIC']


def find_excel():
    candidates = [c for c in glob.glob('*.xlsx') if not os.path.basename(c).startswith('~$')]
    if not candidates:
        print('ERROR: no se encontró ningún archivo .xlsx en la raíz del repositorio.')
        sys.exit(1)
    for c in candidates:
        if os.path.basename(c).lower() == 'asistencia.xlsx':
            return c
    return candidates[0]


def parse_sheet(ws):
    header = [c.value for c in ws[1]]
    month_idx, month_names = [], []
    for i, h in enumerate(header):
        if h and str(h).strip().upper() in MESES_VALIDOS:
            month_idx.append(i)
            month_names.append(str(h).strip().upper())

    people = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[0]
        if name is None or str(name).strip().upper() == 'TOTALES':
            continue
        cargo = row[1] if len(row) > 1 else None
        extra = row[2] if len(row) > 2 else None
        vals = []
        for idx in month_idx:
            v = row[idx] if idx < len(row) else None
            vals.append(str(v).strip().lower() if v else '')
        a, j, n = vals.count('a'), vals.count('j'), vals.count('n')
        total = a + j + n
        pct = (a / total) if total > 0 else None
        people.append({
            'nombre': name, 'cargo': cargo, 'accion': extra,
            'meses': vals, 'a': a, 'j': j, 'n': n, 'pct': pct,
        })
    return people, month_names


def main():
    path = find_excel()
    wb = openpyxl.load_workbook(path, data_only=True)

    out = {'meses_disponibles': [], 'hojas': {}}
    for sheet_name in wb.sheetnames:
        people, month_names = parse_sheet(wb[sheet_name])
        key = sheet_name.strip().lower()
        out['hojas'][key] = people
        if len(month_names) > len(out['meses_disponibles']):
            out['meses_disponibles'] = month_names

    with open('data.json', 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=2)

    total = sum(len(v) for v in out['hojas'].values())
    print(f'OK: procesado "{path}" -> data.json ({total} registros, hojas: {list(out["hojas"].keys())})')


if __name__ == '__main__':
    main()
